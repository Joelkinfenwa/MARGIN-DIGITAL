#!/usr/bin/env python3
"""Builds /e/<slug> pitch pages for e-commerce leads from /home/claude/ecom_leads.json and an 'Ecom pitch' tracker sheet."""
import re, html, json, pathlib, datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

CAL = "https://calendly.com/hello-margindigital/30min"
TEL = "tel:+61449022604"; PHONE = "0449 022 604"; SITE = "https://www.margindigital.com.au"
TRACKER = "/mnt/user-data/outputs/Margin_Digital_Tracker.xlsx"
TODAY = datetime.date.today().strftime("%d %B %Y")

CAT = {
 "supplements": dict(buyer="someone who already takes what you sell and is looking for a better one", searches=["{cat} australia","best {cat} online","{product} buy online","{brand}"], meta="before-and-after and ingredient-led creative, with retargeting on everyone who viewed a product page"),
 "apparel":     dict(buyer="someone scrolling at night who has not heard of you yet", searches=["{cat} australia","australian {cat} brand","{product} online","{brand}"], meta="product-on-body video creative, UGC-style, with dynamic catalogue retargeting of viewed products"),
 "home":        dict(buyer="someone furnishing or gifting who is comparing three tabs", searches=["{product} australia","buy {product} online","{cat} gift","{brand}"], meta="lifestyle and gifting creative, strongest from October through December"),
 "pet":         dict(buyer="an owner who buys the same thing every month", searches=["{product} for dogs","{cat} australia","best {product}","{brand}"], meta="pet-owner video creative and subscription-led retargeting"),
 "beauty":      dict(buyer="someone who found you on a friend's bathroom shelf", searches=["{product} australia","{cat} online","best {product} for","{brand}"], meta="routine and results creative, UGC-led, with retargeting on product viewers"),
 "kids":        dict(buyer="a parent researching late at night", searches=["{product} australia","{cat} online australia","best {product}","{brand}"], meta="parent-to-parent creative and gifting pushes"),
 "outdoor":     dict(buyer="someone planning a weekend and comparing specs", searches=["{product} australia","best {product}","{cat} online","{brand}"], meta="in-use video creative and seasonal pushes"),
 "food":        dict(buyer="a repeat buyer once they have tried it once", searches=["{product} online australia","{cat} delivery","best {product}","{brand}"], meta="first-order offer creative and subscription retargeting"),
}
def catkey(c):
    c=(c or "").lower()
    for k in CAT:
        if k in c: return k
    for k,alts in {"supplements":["health","vitamin","protein"],"apparel":["activewear","clothing","fashion","swim"],"home":["lifestyle","homeware","candle","decor"],"beauty":["skin","cosmetic","hair"],"kids":["baby","toy"],"outdoor":["sport","camping","fitness"],"food":["coffee","tea","snack","drink"]}.items():
        if any(a in c for a in alts): return k
    return "home"

def slug(s): return re.sub(r"[^a-z0-9]+","-",s.lower()).strip("-")

CSS = open("public/p/zaher-associates.html").read().split("<style>")[1].split("</style>")[0]
GTAG = open("public/p/zaher-associates.html").read().split("<head>")[1].split("<meta name=\"viewport\"")[0]

def page(L):
    e=html.escape; name=e(L["store"]); k=catkey(L.get("category")); P=CAT[k]
    product=(L.get("sells") or L.get("category") or "your products").split(",")[0].strip().lower()
    qs=[e(q.format(cat=(L.get("category") or "products").lower(), product=product, brand=L["store"].lower())) for q in P["searches"]]
    size=e(L.get("size_signal") or "")
    return f"""<!doctype html><html lang="en-AU"><head>{GTAG}<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover"><meta name="robots" content="noindex,nofollow">
<title>A growth plan for {name} — Margin Digital</title>
<link rel="icon" href="/favicon.svg" type="image/svg+xml">
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Figtree:wght@500;600;700;800;900&display=swap">
<script defer src="/_vercel/insights/script.js"></script>
<style>{CSS}</style></head><body>
<header class="top"><div class="wrap"><a class="brand" href="{SITE}"><span class="mk"></span>Margin Digital</a><div class="for">Prepared for<br>{name}</div></div></header>
<main>
<section class="hero"><div class="wrap">
  <div class="pill">Private page · prepared {e(TODAY)}</div>
  <h1>A plan to grow {name} with Google Shopping and Meta, run by someone who sells online too.</h1>
  <p class="lead">I put this together before getting in touch so you can see what I would do, what it costs and the guarantee behind it, before deciding whether a conversation is worth your time. Three minutes to read.</p>
  <div class="cta"><a class="btn" href="{CAL}" target="_blank" rel="noopener">Book a free 30-minute call</a><a class="btn ghost" href="{TEL}">Call {PHONE}</a></div>
  <div class="note"><strong>Why you.</strong> {name} has a real product and a store that already converts; {size + ' ' if size else ''}that is the hard part done. The plan below is about buying growth at a return you can see, from a {e(P['buyer'])}.</div>
</div></section>

<section class="alt"><div class="wrap">
  <p class="kicker">Where the next customers are</p>
  <h2>Two channels, one number that matters: what a new customer costs.</h2>
  <div class="q">{''.join(f'<span>{q}</span>' for q in qs)}</div>
  <p class="lead" style="font-size:1.05rem">Google Shopping catches people already searching for what you sell. Meta finds the ones who do not know you yet and brings back the ones who looked and left. Run together, with the same tracking, you see cost per order by channel every week.</p>
</div></section>

<section><div class="wrap">
  <p class="kicker">The plan</p>
  <h2>Three steps, in the order I would do them.</h2>
  <div class="steps">
    <div class="step"><div class="n">1</div><div><h3>Tracking that tells the truth</h3><p>Shopify pixel, Google tag and Meta Conversions API set up server-side so purchases are attributed correctly after iOS changes. No campaign is worth running on numbers you cannot trust.</p></div></div>
    <div class="step"><div class="n">2</div><div><h3>Google Shopping on your best sellers</h3><p>Merchant Centre feed cleaned up, Shopping campaigns on the products with margin to support it, brand search protected, and a target return on ad spend agreed with you up front.</p></div></div>
    <div class="step"><div class="n">3</div><div><h3>Meta for new customers and retargeting</h3><p>{e(P['meta'])}. Creative tested weekly, with you supplying product and me supplying the structure and the testing.</p></div></div>
  </div>
</div></section>

<section class="alt"><div class="wrap">
  <p class="kicker">What it costs</p>
  <h2>Flat fee, not a percentage of your ad spend.</h2>
  <p class="lead" style="font-size:1.05rem">You know your margin and your average order. Compare that to the numbers below.</p>
  <div class="price">
    <div class="plan"><h3>Google Shopping</h3><div class="amt">$799<small>/month</small></div><p>Feed, campaigns, brand protection, weekly optimisation and reporting.</p></div>
    <div class="plan"><h3>Meta Ads</h3><div class="amt">$799<small>/month</small></div><p>Prospecting and retargeting, creative testing, weekly reporting.</p></div>
    <div class="plan hot"><h3>Complete Growth</h3><div class="amt">$1,299<small>/month</small></div><p>Both channels, one tracking setup, cost per order by channel every week.</p></div>
  </div>
  <div class="guar"><b>60</b><div><h3 style="color:#fff">A 60-day setup window. Not happy by day 60? Every dollar of management fees back.</h3><p>Ad spend is paid by you directly to Google and Meta. Your ad accounts, your pixel, your data, in your name from day one. No lock-in after the first 60 days.</p></div></div>
  <p style="margin-top:14px;color:var(--ink-3);font-size:.9rem">AUD, excluding GST. Founding client rates, locked in for as long as you stay.</p>
</div></section>

<section><div class="wrap">
  <p class="kicker">Who is behind this</p>
  <h2>A business owner who runs online stores, not an agency sales team.</h2>
  <p class="lead" style="font-size:1.05rem">I run my own businesses, including online stores on Shopify, and I have hired agencies to run ads for them. More than one let me down: a confident pitch, a long report, and no clear line between what I paid and what came back. Margin Digital is the agency I wanted to hire. You deal with me directly, the prices are on this page, and the accounts are yours.</p>
  <div class="sig">Joel<small>Founder, Margin Digital · {PHONE} · hello@margindigital.com.au</small></div>
</div></section>

<section class="final alt"><div class="wrap">
  <h2>Thirty minutes. I walk you through this plan for {name}, you decide.</h2>
  <div class="cta" style="justify-content:center"><a class="btn" href="{CAL}" target="_blank" rel="noopener">Book a free 30-minute call</a><a class="btn ghost" href="{TEL}">Call {PHONE}</a></div>
  <p class="sub">No obligation. Founding client rates held for the first five stores.</p>
</div></section>
</main>
<footer><div class="wrap">© Margin Digital · Sydney · This page was prepared for {name} and is not indexed or shared.</div></footer>
<div class="sticky"><a class="btn ghost" href="{TEL}">Call</a><a class="btn" href="{CAL}" target="_blank" rel="noopener">Book a free call</a></div>
<script>document.addEventListener('click',function(e){{var a=e.target.closest('a[href*="calendly.com"]');if(a&&window.gtag){{gtag('event','conversion',{{'send_to':'AW-18536302130/KGlTCJmvlKkdELKM5oZF'}});}}var t=e.target.closest('a[href^="tel:"]');if(t&&window.gtag){{gtag('event','conversion',{{'send_to':'AW-18536302130/4WSZCLryz6kdELKM5oZF','value':1.0,'currency':'AUD'}});}}}});</script>
</body></html>"""

def email(L,url):
    subj=f"A growth plan for {L['store']} (three-minute read)"
    body=f"""Hi there,

I am Joel, founder of Margin Digital. I run online stores myself and manage Google Shopping and Meta Ads for Australian brands.

Before reaching out I put together a short plan for {L['store']}: where I would find your next customers, how I would track cost per order by channel, what it costs, and the guarantee behind it. It is a private page, three minutes to read:

{url}

The short version: a flat fee rather than a percentage of spend, your ad accounts and pixel in your name, and if you are not happy by day 60 the management fees are refunded in full.

If it is worth a conversation, you can pick a time here: {CAL}
Or call me directly on {PHONE}.

Kind regards,
Joel
Founder, Margin Digital
{SITE}"""
    return subj, body

leads=json.load(open("/home/claude/ecom_leads.json"))
out=pathlib.Path("public/e"); out.mkdir(exist_ok=True)
recs=[]
for L in leads:
    sl=slug(L["store"]); (out/f"{sl}.html").write_text(page(L))
    url=f"{SITE}/e/{sl}"; subj,body=email(L,url)
    recs.append((L["store"],L.get("category"),L.get("url"),L.get("email"),L.get("instagram"),url,subj,body))

wb=load_workbook(TRACKER)
if "Ecom pitch" in wb.sheetnames: del wb["Ecom pitch"]
ps=wb.create_sheet("Ecom pitch",2)
HDR=PatternFill("solid",fgColor="1F3A5F"); wf=Font(name="Arial",bold=True,color="FFFFFF"); thin=Side(style="thin",color="BFBFBF"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
ps["A1"]="E-commerce pitch pages and emails"; ps["A1"].font=Font(name="Arial",bold=True,size=14)
ps["A2"]="Australian Shopify stores. One private page per store at margindigital.com.au/e/… . Send from hello@margindigital.com.au."; ps["A2"].font=Font(name="Arial",color="595959")
cols=["Store","Category","Store URL","Email","Instagram","Pitch page","Email subject","Email body"]
for i,c in enumerate(cols,1):
    cell=ps.cell(row=4,column=i,value=c); cell.font=wf; cell.fill=HDR; cell.border=bd
for j,r in enumerate(recs,5):
    for i,v in enumerate(r,1):
        c=ps.cell(row=j,column=i,value=v); c.font=Font(name="Arial"); c.border=bd; c.alignment=Alignment(vertical="top",wrap_text=(i==8))
    ps.cell(row=j,column=6).hyperlink=r[5]; ps.cell(row=j,column=6).font=Font(name="Arial",color="0000FF",underline="single")
for i,w in enumerate([30,14,34,30,18,48,44,90],1): ps.column_dimensions[get_column_letter(i)].width=w
ps.freeze_panes="A5"; wb.save(TRACKER)
print(len(recs),"ecom pitch pages")
