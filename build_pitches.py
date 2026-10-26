#!/usr/bin/env python3
"""Builds one private pitch page per emailable lead at public/p/<slug>.html and a Pitch sheet in the tracker."""
import re, html, json, pathlib
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

CAL = "https://calendly.com/hello-margindigital/30min"
TEL = "tel:+61449022604"; PHONE = "0449 022 604"; SITE = "https://www.margindigital.com.au"
TRACKER = "/mnt/user-data/outputs/Margin_Digital_Tracker.xlsx"

# ---- category profiles: customer noun, search phrases ({s} = suburb), plan specifics ----
P = {
 "Legal":        dict(cust="client", noun="clients", q=["solicitor {s}","family lawyer {s}","wills and estates lawyer near me","conveyancing solicitor {s}","lawyer near me open now"], site="a page for each practice area, written so someone searching at 9pm with a problem understands within ten seconds that you handle it and can book or call in one tap", ads="your practice areas, one ad group each, so a search for a family lawyer sees a family-law ad and lands on a family-law page"),
 "Conveyancing": dict(cust="client", noun="settlements", q=["conveyancer {s}","conveyancing {s}","conveyancing quote","property lawyer {s}","cheap conveyancing sydney"], site="a quote request form at the top of the page, your fixed fees stated plainly, and your reviews beside them", ads="conveyancing searches across the Shire and St George, with an instant-quote page as the landing page"),
 "Accounting":   dict(cust="client", noun="clients", q=["accountant {s}","tax accountant near me","small business accountant {s}","bookkeeper {s}","tax return {s}"], site="a page for individuals and a page for business owners, each with a book-a-consultation button and your reviews", ads="searches for a tax accountant and a small business accountant in your area, strongest from July to October"),
 "Finance":      dict(cust="client", noun="clients", q=["financial advisor {s}","financial planner near me","retirement planning {s}","superannuation advice {s}","mortgage broker {s}"], site="a clear first-meeting offer, your qualifications and reviews up front, and a booking calendar on the page", ads="searches for a financial adviser or planner in the Shire, landing on a page built around the first meeting"),
 "Mortgage broker": dict(cust="client", noun="clients", q=["mortgage broker {s}","home loan broker near me","refinance {s}","first home buyer broker {s}","best mortgage broker sutherland shire"], site="a borrowing-power or refinance check as the main call to action, with your reviews next to it", ads="refinance and first-home-buyer searches in your area, with a short form as the landing page"),
 "Physio":       dict(cust="patient", noun="patients", q=["physio {s}","physiotherapist near me","sports physio {s}","back pain physio {s}","physio open saturday {s}"], site="online booking in the first screen on a phone, a page per condition you treat, and your reviews where people decide", ads="searches for a physio in your suburb and the ones around it, with same-week availability in the ad"),
 "Chiro":        dict(cust="patient", noun="patients", q=["chiropractor {s}","chiro near me","back pain chiropractor {s}","neck pain treatment {s}","chiropractor open saturday"], site="online booking in the first screen on a phone, a page per condition, and your reviews where people decide", ads="chiropractor searches in your suburb and surrounds, landing on a booking-first page"),
 "Massage":      dict(cust="client", noun="bookings", q=["remedial massage {s}","massage near me","sports massage {s}","deep tissue massage {s}","remedial massage health fund rebate"], site="online booking in the first screen, health-fund rebates stated clearly, and your reviews beside the book button", ads="remedial and sports massage searches in your area, which are high-intent and often same-day"),
 "Day spa":      dict(cust="client", noun="bookings", q=["day spa {s}","facial {s}","massage {s}","spa packages sutherland shire","gift voucher day spa"], site="treatment menu with prices, online booking at the top, and gift vouchers as a second call to action", ads="day spa and facial searches across the Shire, with seasonal push on vouchers before Mother's Day and Christmas"),
 "Cosmetic":     dict(cust="client", noun="clients", q=["cosmetic clinic {s}","anti wrinkle injections {s}","dermal fillers near me","skin clinic {s}","laser hair removal {s}"], site="a page per treatment with pricing or a price guide, before-and-after where appropriate, and booking in one tap", ads="treatment-level searches in your area, each landing on its own page, which is what keeps the cost per booking down"),
 "Dental":       dict(cust="patient", noun="patients", q=["dentist {s}","emergency dentist near me","teeth whitening {s}","dental implants {s}","dentist open saturday {s}"], site="new-patient offer and online booking in the first screen, a page per treatment, and your reviews up front", ads="emergency and new-patient searches in your suburb, which convert fastest"),
 "Vet":          dict(cust="client", noun="new clients", q=["vet {s}","vet near me","emergency vet {s}","puppy vaccination {s}","cat vet {s}"], site="online booking and click-to-call in the first screen, a new-client page, and your reviews where people decide", ads="vet near me and new-pet searches in your suburb and surrounds"),
 "Mechanic":     dict(cust="customer", noun="bookings", q=["mechanic {s}","car service {s}","pink slip {s}","brake repair near me","logbook service {s}"], site="a book-a-service form at the top, services and indicative prices, and your reviews beside the form", ads="car service, pink slip and brake searches in your suburb, which are high-intent and often same-week"),
 "Renovations":  dict(cust="client", noun="quote requests", q=["kitchen renovation {s}","bathroom renovation {s}","kitchen renovations sutherland shire","bathroom renovators near me","renovation cost sydney"], site="a get-a-quote form at the top, your work shown large, and a price guide so the enquiries are qualified", ads="kitchen and bathroom renovation searches across the Shire and St George, landing on a quote page"),
 "Building":     dict(cust="client", noun="quote requests", q=["builder {s}","home builder sutherland shire","house extension {s}","knockdown rebuild {s}","custom home builder sydney"], site="a get-a-quote form at the top, your projects shown large, and a clear process section", ads="extension, knockdown-rebuild and custom-home searches across southern Sydney"),
 "Pools":        dict(cust="client", noun="quote requests", q=["pool builder {s}","concrete pools sutherland shire","pool renovation {s}","pool builders near me","plunge pool sydney"], site="a get-a-quote form at the top, your pools shown large, and a pool-type page for each kind you build", ads="pool builder and pool renovation searches across southern Sydney, strongest August to December"),
 "Solar":        dict(cust="customer", noun="quote requests", q=["solar installer {s}","solar panels {s}","solar battery {s}","solar quote sydney","solar rebate nsw"], site="a quote form at the top, rebates explained, and your installs shown with suburb names", ads="solar and battery searches across southern Sydney, landing on a quote page"),
 "Real estate":  dict(cust="vendor", noun="appraisal requests", q=["real estate agent {s}","sell my house {s}","property appraisal {s}","best real estate agent {s}","rental property management {s}"], site="a free-appraisal form at the top, recent results by suburb, and your reviews beside the form", ads="appraisal and 'sell my house' searches in your core suburbs, landing on an appraisal page"),
 "Swim school":  dict(cust="family", noun="enrolments", q=["swim school {s}","swimming lessons {s}","learn to swim near me","baby swimming lessons {s}","adult swimming lessons {s}"], site="an enrol-or-book-a-trial form at the top, timetable and prices, and your reviews beside it", ads="swimming-lesson searches in your area, strongest in the month before each term"),
 "Gym":          dict(cust="member", noun="members", q=["gym {s}","24 hour gym near me","personal trainer {s}","group fitness {s}","gym membership {s}"], site="a free-trial or intro-offer form at the top, memberships and timetable, and your reviews beside it", ads="gym and personal trainer searches in your suburb, landing on a trial-offer page"),
 "Pilates":      dict(cust="client", noun="members", q=["pilates {s}","reformer pilates near me","pilates studio {s}","clinical pilates {s}","pilates intro offer"], site="an intro-offer form at the top, timetable and pricing, and your reviews beside it", ads="reformer and clinical pilates searches in your area, landing on an intro-offer page"),
 "Therapy":      dict(cust="client", noun="clients", q=["therapy {s}","counselling near me","psychologist {s}","anxiety therapy {s}","couples counselling {s}"], site="a first-session booking at the top, a page per concern you help with, and clear fees", ads="counselling and therapy searches in your area, landing on a first-session page"),
 "Psychology":   dict(cust="client", noun="clients", q=["psychologist {s}","psychologist near me","child psychologist {s}","anxiety psychologist {s}","bulk billing psychologist {s}"], site="a first-session booking at the top, a page per concern you help with, and clear fees and rebates", ads="psychologist searches in your area, landing on a first-session page"),
 "Podiatry":     dict(cust="patient", noun="patients", q=["podiatrist {s}","podiatrist near me","ingrown toenail {s}","heel pain treatment {s}","orthotics {s}"], site="online booking in the first screen, a page per condition, and your reviews where people decide", ads="podiatrist and condition searches in your suburb and surrounds"),
 "Studio":       dict(cust="client", noun="members", q=["pilates {s}","yoga {s}","reformer pilates near me","barre {s}","studio intro offer {s}"], site="an intro-offer form at the top, timetable and pricing, and your reviews beside it", ads="studio and intro-offer searches in your area"),
 "Generic":      dict(cust="client", noun="enquiries", q=["{biz} {s}","{cat} near me","{cat} {s}","best {cat} sutherland shire","{cat} reviews"], site="one clear call to action in the first screen on a phone, a page per service, and your reviews where people decide", ads="searches for your services in your suburb and the ones around it, with tracking on every call and form"),
}
WHAT = {"Legal":"a lawyer","Conveyancing":"a conveyancer","Accounting":"an accountant","Finance":"a financial adviser","Mortgage broker":"a mortgage broker","Physio":"a physio","Chiro":"a chiropractor","Massage":"a remedial massage therapist","Day spa":"a day spa","Cosmetic":"a cosmetic clinic","Dental":"a dentist","Vet":"a vet","Mechanic":"a mechanic","Renovations":"a renovator","Building":"a builder","Pools":"a pool builder","Solar":"a solar installer","Real estate":"an agent","Swim school":"swimming lessons","Gym":"a gym","Pilates":"a pilates studio","Therapy":"a therapist","Psychology":"a psychologist","Podiatry":"a podiatrist","Studio":"a studio","Generic":"a business like yours"}
# manual category fixes for rows with no category
FIX = {"Zaher & Associates":"Accounting","Engadine Health":"Gym","Powe Financial Advice Pty Ltd":"Finance","Top Notch":"Conveyancing","Under The Bar":"Gym","Fitwise PTY LTD":"Physio","Colin Daley Quinn":"Legal","Alasana":"Studio","Kids Aquatic Survival School":"Swim school","Everingham & Turner":"Generic","Prosperity Associates":"Accounting","GJ Private Wealth (Sutherland Shire)":"Finance","ND Renovations":"Renovations","James Hayes - Financial Planner":"Finance","Tailored Homes":"Building","Ledgersmith Sutherland":"Accounting","AGS Financial Group Sutherland":"Finance"}
NICE = {"Zaher & Associates":"Zaher & Associates","Top Notch":"Top Notch Conveyancing","Fitwise PTY LTD":"Fitwise","GJ Private Wealth (Sutherland Shire)":"GJ Private Wealth","James Hayes - Financial Planner":"James Hayes Financial Planning","Ledgersmith Sutherland":"Ledgersmith","AGS Financial Group Sutherland":"AGS Financial Group","FIRM RELIEVE SPORTS & REMEDIAL MASSAGE":"Firm Relieve","Your Physiotherapy @ Rockdale Pty Ltd":"Your Physiotherapy Rockdale","LJS Conveyancing Services Pty Ltd":"LJS Conveyancing","Move Therapy - Rockdale Physiotherapy & Exercise Physiology":"Move Therapy","Sullivan Dewing Chartered Accountants & Business Builders":"Sullivan Dewing","Rick Pender's Auto Service":"Rick Pender's Auto Service","Powe Financial Advice Pty Ltd":"Powe Financial Advice","HT Wills Real Estate St George":"HT Wills Real Estate"}

def slug(s): return re.sub(r"[^a-z0-9]+","-",s.lower()).strip("-")
def nice(n):
    n = NICE.get(n, n)
    n = re.sub(r"\s+(pty\s*ltd|pty|ltd)\.?$", "", n, flags=re.I)
    return n if n.isupper() is False else n.title()

CSS = """
:root{--bg:#fff;--surface:#F3F6FA;--ink:#0B1E3F;--ink-2:#3E4C66;--ink-3:#7A8699;--line:#D5DCE8;--accent:#2449E6;--mark:#D4F23A;--ok:#1C9B5E;--font:"Figtree","Helvetica Neue",Arial,sans-serif;color-scheme:light}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font-family:var(--font);font-size:18px;line-height:1.55;font-weight:500;-webkit-font-smoothing:antialiased}
em,i{font-style:normal}a{color:inherit}p{margin:0}
h1,h2,h3{margin:0;line-height:1.08;letter-spacing:-.02em;font-weight:800;text-wrap:balance}
h1{font-size:clamp(2rem,4.8vw,3.4rem);font-weight:900;letter-spacing:-.03em}h2{font-size:clamp(1.5rem,3vw,2.1rem)}h3{font-size:1.15rem}
.wrap{max-width:860px;margin:0 auto;padding:0 max(20px,env(safe-area-inset-left,0px))}
.btn{display:inline-flex;align-items:center;justify-content:center;font:inherit;font-weight:800;font-size:1.05rem;padding:16px 26px;border-radius:10px;text-decoration:none;background:var(--accent);color:#fff;box-shadow:0 8px 24px rgba(36,73,230,.25)}
.btn:hover{background:#1B3BC4}.btn.ghost{background:#fff;color:var(--ink);border:2px solid var(--line);box-shadow:none}
.top{border-bottom:1px solid var(--line);position:sticky;top:0;z-index:10;background:rgba(255,255,255,.95);backdrop-filter:blur(8px)}
.top .wrap{display:flex;align-items:center;justify-content:space-between;padding-block:12px;gap:12px}
.brand{font-weight:900;font-size:1.15rem;letter-spacing:-.03em;text-decoration:none;display:flex;align-items:center;gap:9px}.brand .mk{width:13px;height:13px;background:var(--accent);border-radius:3px;transform:rotate(45deg)}
.top .for{color:var(--ink-3);font-weight:700;font-size:.9rem;text-align:right}
section{padding-block:clamp(40px,6vw,64px)}
.kicker{font-weight:700;color:var(--accent);margin-bottom:12px}
.lead{font-size:clamp(1.1rem,1.6vw,1.3rem);color:var(--ink-2);margin-top:16px}
.pill{display:inline-flex;align-items:center;gap:8px;background:var(--mark);color:var(--ink);font-weight:800;font-size:.88rem;padding:7px 13px;border-radius:999px;margin-bottom:18px}
.hero{padding-top:clamp(36px,5vw,64px)}
.cta{margin-top:24px;display:flex;flex-wrap:wrap;gap:10px}
.note{margin-top:22px;padding:18px 20px;border-left:4px solid var(--accent);background:var(--surface);border-radius:0 12px 12px 0;color:var(--ink-2)}
.facts{display:grid;grid-template-columns:repeat(3,1fr);gap:14px;margin-top:22px}
@media (max-width:640px){.facts{grid-template-columns:1fr 1fr}}
.fact{background:var(--surface);border-radius:14px;padding:18px}.fact b{display:block;font-size:1.7rem;font-weight:900;letter-spacing:-.03em;line-height:1}.fact span{display:block;margin-top:6px;color:var(--ink-2);font-size:.95rem;font-weight:600}
.q{display:flex;flex-wrap:wrap;gap:8px;margin-top:16px}.q span{background:#fff;border:1px solid var(--line);border-radius:999px;padding:7px 14px;font-weight:700;font-size:.95rem}
.alt{background:var(--surface)}
.steps{display:grid;gap:14px;margin-top:22px}
.step{background:#fff;border:1px solid var(--line);border-radius:16px;padding:22px;display:grid;grid-template-columns:44px 1fr;gap:16px}
.step .n{width:36px;height:36px;border-radius:10px;background:var(--mark);font-weight:900;display:grid;place-content:center}
.step p{color:var(--ink-2);margin-top:6px;font-size:1rem}
.price{display:grid;grid-template-columns:repeat(3,1fr);gap:14px;margin-top:22px}
@media (max-width:700px){.price{grid-template-columns:1fr}}
.plan{border:1px solid var(--line);border-radius:16px;padding:22px}.plan.hot{background:var(--ink);color:#fff;border-color:var(--ink)}
.plan .amt{font-size:2rem;font-weight:900;letter-spacing:-.03em;line-height:1;margin-top:8px}.plan .amt small{font-size:.9rem;font-weight:700;letter-spacing:0;color:var(--ink-3);margin-left:6px}.plan.hot .amt small{color:#C9D3E6}
.plan p{margin-top:10px;font-size:.98rem;color:var(--ink-2)}.plan.hot p{color:#C9D3E6}
.guar{background:var(--ink);color:#fff;border-radius:20px;padding:28px;margin-top:24px;display:grid;grid-template-columns:auto 1fr;gap:22px;align-items:center}
@media (max-width:640px){.guar{grid-template-columns:1fr}}
.guar b{font-size:2.6rem;font-weight:900;letter-spacing:-.04em;line-height:1;color:var(--mark)}.guar p{color:#C9D3E6}
.sig{margin-top:18px;font-weight:800}.sig small{display:block;font-weight:600;color:var(--ink-3)}
.final{text-align:center}.final .sub{margin-top:14px;color:var(--ink-3);font-weight:600}
footer{border-top:1px solid var(--line);padding:24px 0 100px;color:var(--ink-3);font-size:.88rem}
.sticky{display:none}@media (max-width:860px){.sticky{display:flex;gap:10px;position:fixed;left:0;right:0;bottom:0;z-index:10;padding:10px 16px calc(10px + env(safe-area-inset-bottom,0px));background:rgba(255,255,255,.96);backdrop-filter:blur(8px);border-top:1px solid var(--line)}.sticky .btn{flex:1;padding:14px 10px;font-size:1rem}.sticky .btn.ghost{flex:0 0 auto}}
"""

GTAG = """<script>
  window.dataLayer = window.dataLayer || [];function gtag(){dataLayer.push(arguments);}gtag('js', new Date());gtag('config', 'AW-18536302130');
  (function(){var done=false;function load(){if(done)return;done=true;var s=document.createElement('script');s.async=true;s.src='https://www.googletagmanager.com/gtag/js?id=AW-18536302130';document.head.appendChild(s);}
  window.addEventListener('load',function(){setTimeout(load,1500);});['pointerdown','touchstart','keydown','scroll'].forEach(function(ev){window.addEventListener(ev,load,{once:true,passive:true});});})();
</script>"""

def page(biz, cat, suburb, rating, reviews, prof):
    e = html.escape
    name = e(biz); s = e(suburb or "your area")
    qs = [e(q.format(s=suburb or "near me", biz=biz, cat=(cat or "service").lower())) for q in prof["q"]]
    rev_line = f"{int(reviews)} Google reviews at {rating:g} stars" if reviews and rating else "a strong reputation"
    rev_fact = f'<div class="fact"><b>{int(reviews)}</b><span>Google reviews</span></div><div class="fact"><b>{rating:g}★</b><span>average rating</span></div>' if reviews and rating else ""
    return f"""<!doctype html><html lang="en-AU"><head><meta charset="utf-8">{GTAG}
<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover"><meta name="robots" content="noindex,nofollow">
<title>A plan for {name} — Margin Digital</title>
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Figtree:wght@500;600;700;800;900&display=swap">
<style>{CSS}</style></head><body>
<header class="top"><div class="wrap"><a class="brand" href="{SITE}"><span class="mk"></span>Margin Digital</a><div class="for">Prepared for<br>{name}</div></div></header>
<main>
<section class="hero"><div class="wrap">
  <div class="pill">Private page · prepared {e(TODAY)}</div>
  <h1>A plan to turn {name}'s reputation into more {e(prof['noun'])} in {s}.</h1>
  <p class="lead">I put this together before getting in touch, so you can see exactly what I would do and what it costs before deciding whether a conversation is worth your time. It takes about three minutes to read.</p>
  <div class="cta"><a class="btn" href="{CAL}" target="_blank" rel="noopener">Book a free 30-minute call</a><a class="btn ghost" href="{TEL}">Call {PHONE}</a></div>
  <div class="note"><strong>Why you.</strong> {name} has {rev_line}. That is the hardest part of marketing already done. The plan below is about putting it in front of people in {s} at the moment they are searching for {e(WHAT.get(cat,'a business like yours'))}.</div>
  <div class="facts">{rev_fact}<div class="fact"><b>{s}</b><span>home suburb and surrounds</span></div></div>
</div></section>

<section class="alt"><div class="wrap">
  <p class="kicker">What people in {s} are typing into Google</p>
  <h2>Searches like these happen every week. The plan is to be the business they choose.</h2>
  <div class="q">{''.join(f'<span>{q}</span>' for q in qs)}</div>
  <p class="lead" style="font-size:1.05rem">Each one of these is a person with a need and a phone in their hand. Being the clear, easy choice when they search is where the next {e(prof['noun'])} come from.</p>
</div></section>

<section><div class="wrap">
  <p class="kicker">The plan</p>
  <h2>Three steps, in the order I would do them.</h2>
  <div class="steps">
    <div class="step"><div class="n">1</div><div><h3>A website built to convert, not just to exist</h3><p>For {name} that means {e(prof['site'])}. Written for your business, built mobile-first, with calls and enquiries wired to your phone. Fixed $2,500, once.</p></div></div>
    <div class="step"><div class="n">2</div><div><h3>Your Google listing and reviews put to work</h3><p>Your Google Business Profile tuned so the map result shows the right services, hours and photos, with your reviews doing the convincing. Included with the website.</p></div></div>
    <div class="step"><div class="n">3</div><div><h3>Google Ads on the searches above</h3><p>Ads on {e(prof['ads'])}. Every call and form tracked, so each month you see what was spent, how many enquiries came in and what each one cost. Flat $799 a month, not a percentage of spend.</p></div></div>
  </div>
</div></section>

<section class="alt"><div class="wrap">
  <p class="kicker">What it costs</p>
  <h2>Fixed prices, published upfront. Founding client rates, locked in.</h2>
  <p class="lead" style="font-size:1.05rem">You know better than anyone what a new {e(prof['cust'])} is worth to {name}. Compare that to the numbers below.</p>
  <div class="price">
    <div class="plan"><h3>Website</h3><div class="amt">$2,500<small>once</small></div><p>Design, copy, build, Google Business Profile. No monthly fee for the site.</p></div>
    <div class="plan"><h3>Google Ads</h3><div class="amt">$799<small>/month</small></div><p>Flat fee. Your own account, your data, no lock-in. Ad spend paid by you directly to Google.</p></div>
    <div class="plan hot"><h3>Complete Growth</h3><div class="amt">$1,299<small>/month</small></div><p>Google Ads and Meta Ads together, one report, cost per enquiry by channel.</p></div>
  </div>
  <div class="guar"><b>100%</b><div><h3 style="color:#fff">Not happy at any point? Every dollar back.</h3><p>No conditions, no minimum term. If you decide Margin Digital is not for you, say so and you are refunded in full. You keep the site, the listing and the ad account.</p></div></div>
  <p style="margin-top:14px;color:var(--ink-3);font-size:.9rem">AUD, excluding GST. Hosting and third-party subscriptions are separate and itemised up front.</p>
</div></section>

<section><div class="wrap">
  <p class="kicker">Who is behind this</p>
  <h2>A business owner, not an agency sales team.</h2>
  <p class="lead" style="font-size:1.05rem">I run my own business and I have hired agencies to do what I now do. More than one let me down: a confident pitch, a long report, and no clear line between what I paid and what I got back. Margin Digital is the agency I wanted to hire. You deal with me directly, the prices are on this page, every account is in your name, and if you are not happy at any point you get every dollar back.</p>
  <div class="sig">Joel<small>Founder, Margin Digital · {PHONE} · hello@margindigital.com.au</small></div>
</div></section>

<section class="final alt"><div class="wrap">
  <h2>Thirty minutes. I walk you through this plan for {name}, you decide.</h2>
  <div class="cta" style="justify-content:center"><a class="btn" href="{CAL}" target="_blank" rel="noopener">Book a free 30-minute call</a><a class="btn ghost" href="{TEL}">Call {PHONE}</a></div>
  <p class="sub">No obligation. Founding client rates held for the first five businesses.</p>
</div></section>
</main>
<footer><div class="wrap">© Margin Digital · Sydney · This page was prepared for {name} and is not indexed or shared.</div></footer>
<div class="sticky"><a class="btn ghost" href="{TEL}">Call</a><a class="btn" href="{CAL}" target="_blank" rel="noopener">Book a free call</a></div>
<script>document.addEventListener('click',function(e){{var a=e.target.closest('a[href*="calendly.com"]');if(a&&window.gtag){{gtag('event','conversion',{{'send_to':'AW-18536302130/KGlTCJmvlKkdELKM5oZF'}});}}var t=e.target.closest('a[href^="tel:"]');if(t&&window.gtag){{gtag('event','conversion',{{'send_to':'AW-18536302130/4WSZCLryz6kdELKM5oZF','value':1.0,'currency':'AUD'}});}}}});</script>
</body></html>"""

def email(biz, suburb, reviews, prof, url, what):
    first = "there"
    subj = f"A plan for {biz} (three-minute read)"
    body = f"""Hi {first},

I am Joel, founder of Margin Digital, a small Sydney agency that builds websites and runs Google Ads for local businesses.

Before reaching out I put together a short plan for {biz}: what people in {suburb or 'your area'} are searching for, what I would build, what it would cost, and the guarantee that sits behind it. It is a private page, three minutes to read:

{url}

The short version: {biz} has {int(reviews) if reviews else 'a strong set of'} Google reviews, which is the hardest part already done. The plan is about putting that in front of the people searching for {what} at the moment they need one. Fixed prices, and if you are not happy at any point, every dollar back.

If it is worth a conversation, you can pick a time here: {CAL}
Or call me directly on {PHONE}.

Kind regards,
Joel
Founder, Margin Digital
{SITE}"""
    return subj, body

import datetime
TODAY = datetime.date.today().strftime("%d %B %Y")

wb = load_workbook(TRACKER)
ws = wb["Leads"]
rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0] and r[6]]
out = pathlib.Path("public/p"); out.mkdir(exist_ok=True)
recs = []
for r in rows:
    raw, cat, suburb, rating, reviews, phone, em, site = r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7]
    cat = FIX.get(raw, cat) or "Generic"
    prof = P.get(cat, P["Generic"])
    biz = nice(raw)
    sl = slug(biz)
    (out / f"{sl}.html").write_text(page(biz, cat, suburb, rating, reviews, prof))
    url = f"{SITE}/p/{sl}"
    subj, body = email(biz, suburb, reviews, prof, url, WHAT.get(cat,'a business like yours'))
    recs.append((biz, cat, suburb, em, url, subj, body))

# Pitch sheet
if "Pitch" in wb.sheetnames: del wb["Pitch"]
ps = wb.create_sheet("Pitch", 1)
HDR = PatternFill("solid", fgColor="1F3A5F"); wb_f = Font(name="Arial", bold=True, color="FFFFFF"); thin = Side(style="thin", color="BFBFBF"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
ps["A1"] = "Personalised pitch pages and emails"; ps["A1"].font = Font(name="Arial", bold=True, size=14)
ps["A2"] = "One private page per business at margindigital.com.au/p/… . Copy the subject and body into Gmail, send from hello@margindigital.com.au, then set Status on the Leads tab."; ps["A2"].font = Font(name="Arial", color="595959")
cols = ["Business", "Category", "Suburb", "Email", "Pitch page", "Email subject", "Email body"]
for i, c in enumerate(cols, 1):
    cell = ps.cell(row=4, column=i, value=c); cell.font = wb_f; cell.fill = HDR; cell.border = bd; cell.alignment = Alignment(vertical="center")
for j, rec in enumerate(recs, 5):
    for i, v in enumerate(rec, 1):
        c = ps.cell(row=j, column=i, value=v); c.font = Font(name="Arial"); c.border = bd; c.alignment = Alignment(vertical="top", wrap_text=(i == 7))
    ps.cell(row=j, column=5).hyperlink = rec[4]; ps.cell(row=j, column=5).font = Font(name="Arial", color="0000FF", underline="single")
for i, w in enumerate([34, 14, 14, 30, 48, 42, 90], 1): ps.column_dimensions[get_column_letter(i)].width = w
ps.freeze_panes = "A5"
wb.save(TRACKER)
json.dump([dict(zip(cols, r)) for r in recs], open("public/p/_index.json", "w"), indent=1)
print(len(recs), "pitch pages")
